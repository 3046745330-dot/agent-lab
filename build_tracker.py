# -*- coding: utf-8 -*-
"""生成投递记录表（45 天求职作战用）。

列结构按「驱动行动」设计，不是简单罗列：
  - 状态用数据验证下拉，避免手打不一致
  - 加「下一步」和「下次跟进日」两列 —— 投递最大的失败是投完就忘
  - 加「复盘」列 —— 面试完立刻填，这是把「不会」变「会」的唯一手段

用法：
  python build_tracker.py [输出路径]
"""
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

OUT = (
    sys.argv[1]
    if len(sys.argv) > 1
    else r"C:\Users\yyy\Documents\deepseek-harness\default-workspace\agent-lab\投递记录表.xlsx"
)

HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
HEADER_FONT = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
BODY_FONT = Font(name="微软雅黑", size=10)
NOTE_FONT = Font(name="微软雅黑", size=9, italic=True, color="666666")
THIN = Side(style="thin", color="BBBBBB")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")

wb = Workbook()

# ============================================================
# Sheet 1: 投递记录
# ============================================================
ws = wb.active
ws.title = "投递记录"

COLUMNS = [
    ("序号", 6),
    ("投递日期", 12),
    ("公司", 20),
    ("岗位", 22),
    ("渠道", 12),
    ("方向", 10),
    ("状态", 12),
    ("下一步", 18),
    ("下次跟进日", 12),
    ("面试时间", 16),
    ("面试复盘", 40),
    ("备注", 24),
]

for i, (name, width) in enumerate(COLUMNS, start=1):
    c = ws.cell(row=1, column=i, value=name)
    c.fill = HEADER_FILL
    c.font = HEADER_FONT
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c.border = BORDER
    ws.column_dimensions[get_column_letter(i)].width = width
ws.row_dimensions[1].height = 24
ws.freeze_panes = "A2"

# 预填 60 行（45 天投 60-100 家是合理量）
for r in range(2, 62):
    ws.cell(row=r, column=1, value=r - 1).font = BODY_FONT
    for i in range(1, len(COLUMNS) + 1):
        cell = ws.cell(row=r, column=i)
        cell.font = BODY_FONT
        cell.border = BORDER
        cell.alignment = WRAP
    ws.row_dimensions[r].height = 20

# --- 数据验证下拉 ---
def add_dv(formula, cols, title):
    dv = DataValidation(type="list", formula1=formula, allow_blank=True, showDropDown=False)
    dv.promptTitle = title
    dv.prompt = f"选择{title}"
    ws.add_data_validation(dv)
    for col in cols:
        dv.add(f"{col}2:{col}61")
    return dv

add_dv('"BOSS直聘,实习僧,牛客网,官网,V2EX,内推,其他"', ["E"], "渠道")
add_dv('"AI Agent,Python后端,全栈/前端,其他"', ["F"], "方向")
add_dv(
    '"已投递,已读,HR已回,约面中,一面,二面,三面,HR面,已发offer,已拒,被拒,无响应"',
    ["G"],
    "状态",
)

# --- 状态条件格式（一眼看出进展）---
from openpyxl.formatting.rule import CellIsRule

GREEN = PatternFill("solid", fgColor="D4EDDA")
YELLOW = PatternFill("solid", fgColor="FFF3CD")
RED = PatternFill("solid", fgColor="F8D7DA")
BLUE = PatternFill("solid", fgColor="D1ECF1")

for text, fill in [
    ("已发offer", GREEN),
    ("已拒", RED),
    ("被拒", RED),
    ("无响应", RED),
    ("约面中", YELLOW),
    ("一面", YELLOW),
    ("二面", YELLOW),
    ("三面", YELLOW),
    ("HR面", YELLOW),
    ("HR已回", BLUE),
    ("已读", BLUE),
]:
    ws.conditional_formatting.add(
        f"G2:G61",
        CellIsRule(operator="equal", formula=[f'"{text}"'], fill=fill),
    )

# --- 示例行（教怎么填，用户可删）---
example = [
    1, "2026-10-17", "示例公司（可删）", "AI应用开发实习生", "BOSS直聘",
    "AI Agent", "已投递", "等回复，3天后未回则跟进", "2026-10-20",
    "", "", "打招呼话术已发，附 GitHub 链接",
]
for i, v in enumerate(example, start=1):
    c = ws.cell(row=2, column=i, value=v)
    c.font = Font(name="微软雅黑", size=10, color="999999", italic=True)

# ============================================================
# Sheet 2: 面试复盘
# ============================================================
ws2 = wb.create_sheet("面试复盘")
COLS2 = [
    ("日期", 12),
    ("公司", 18),
    ("轮次", 10),
    ("时长(分)", 10),
    ("结果", 10),
    ("被问到的题", 45),
    ("答得不好的", 45),
    ("应该怎么答", 45),
    ("面试官反馈", 30),
    ("下次改进", 30),
]
for i, (name, width) in enumerate(COLS2, start=1):
    c = ws2.cell(row=1, column=i, value=name)
    c.fill = HEADER_FILL
    c.font = HEADER_FONT
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c.border = BORDER
    ws2.column_dimensions[get_column_letter(i)].width = width
ws2.row_dimensions[1].height = 24
ws2.freeze_panes = "A2"

for r in range(2, 42):
    for i in range(1, len(COLS2) + 1):
        cell = ws2.cell(row=r, column=i)
        cell.font = BODY_FONT
        cell.border = BORDER
        cell.alignment = WRAP
    ws2.row_dimensions[r].height = 30

dv2 = DataValidation(
    type="list",
    formula1='"待定,通过,挂,进下一轮"',
    allow_blank=True,
    showDropDown=False,
)
ws2.add_data_validation(dv2)
dv2.add("E2:E41")

# ============================================================
# Sheet 3: 八股进度
# ============================================================
ws3 = wb.create_sheet("八股进度")
COLS3 = [("优先级", 10), ("模块", 20), ("具体内容", 50), ("预估(h)", 10), ("状态", 12), ("完成日期", 12)]
for i, (name, width) in enumerate(COLS3, start=1):
    c = ws3.cell(row=1, column=i, value=name)
    c.fill = HEADER_FILL
    c.font = HEADER_FONT
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c.border = BORDER
    ws3.column_dimensions[get_column_letter(i)].width = width
ws3.row_dimensions[1].height = 24
ws3.freeze_panes = "A2"

TOPICS = [
    ("P0", "Python 基础", "GIL 是什么、为什么存在、对多线程的影响", 2),
    ("P0", "Python 基础", "装饰器原理与常见用法（含带参数的装饰器）", 2),
    ("P0", "Python 基础", "生成器 / 迭代器 / yield 的执行流程", 2),
    ("P0", "Python 基础", "深浅拷贝的区别与实现原理", 1),
    ("P0", "Python 基础", "上下文管理器 with 与 __enter__/__exit__", 1),
    ("P0", "HTTP", "GET/POST/PUT/DELETE 语义与幂等性", 1),
    ("P0", "HTTP", "状态码：301/302/303/307 区别、4xx/5xx 常见码", 2),
    ("P0", "HTTP", "HTTP/1.1 vs 2 vs 3 的核心差异", 2),
    ("P0", "WebSocket", "握手过程、与 HTTP 长连接的区别、心跳与重连", 2),
    ("P0", "数据库", "索引原理（B+ 树）、聚簇索引 vs 二级索引", 3),
    ("P0", "数据库", "联合索引最左前缀原则", 2),
    ("P0", "数据库", "深分页为什么慢、三种优化方案", 2),
    ("P0", "数据库", "事务 ACID、隔离级别、MVCC", 3),
    ("P1", "消息队列", "为什么要用 MQ、Kafka/RabbitMQ/RocketMQ 选型对比", 2),
    ("P1", "消息队列", "分区与消费者关系、分区键与顺序性", 2),
    ("P1", "消息队列", "rebalance 触发条件与影响", 1),
    ("P1", "消息队列", "auto ACK vs manual ACK、幂等、死信队列", 2),
    ("P1", "LLM 应用", "token / 上下文窗口 / 无状态 API 的含义", 2),
    ("P1", "LLM 应用", "Function Calling 的完整流程与 tools 字段", 2),
    ("P1", "LLM 应用", "RAG 全流程：切分、向量化、检索、重排", 3),
    ("P1", "LLM 应用", "Agent 主循环（ReAct）与失败处理", 2),
    ("P2", "Redis", "常用数据结构与典型场景", 2),
    ("P2", "Redis", "缓存穿透 / 击穿 / 雪崩的区别与对策", 2),
    ("P2", "FastAPI", "依赖注入、异步、与 Flask 的差异", 2),
]

r = 2
for pri, mod, content, hours in TOPICS:
    ws3.cell(row=r, column=1, value=pri)
    ws3.cell(row=r, column=2, value=mod)
    ws3.cell(row=r, column=3, value=content)
    ws3.cell(row=r, column=4, value=hours)
    ws3.cell(row=r, column=5, value="未开始")
    for i in range(1, 7):
        cell = ws3.cell(row=r, column=i)
        cell.font = BODY_FONT
        cell.border = BORDER
        cell.alignment = WRAP
    r += 1

total_row = r
ws3.cell(row=total_row, column=3, value="合计").font = Font(name="微软雅黑", size=10, bold=True)
ws3.cell(row=total_row, column=4, value=f"=SUM(D2:D{r-1})").font = Font(
    name="微软雅黑", size=10, bold=True
)

dv3 = DataValidation(type="list", formula1='"未开始,学习中,已掌握,待复习"', allow_blank=True, showDropDown=False)
ws3.add_data_validation(dv3)
dv3.add(f"E2:E{r-1}")

# ============================================================
# Sheet 4: 说明
# ============================================================
ws4 = wb.create_sheet("使用说明")
ws4.column_dimensions["A"].width = 100
NOTES = [
    "投递记录表 · 使用说明",
    "",
    "【为什么要有这张表】",
    "投递最大的失败不是被拒，是「投完就忘」——不知道自己投了哪些、",
    "哪些该跟进了、哪些面完了没复盘。这张表就是防这个。",
    "",
    "【每天怎么用】",
    "1. 投完一家，立刻在「投递记录」填一行（30 秒的事，别攒着）",
    "2. 「下一步」和「下次跟进日」必须填——这是这张表的价值所在",
    "3. 收到面试邀请，改状态为「约面中」，填面试时间",
    "4. 面完当天，去「面试复盘」页填一行，越详细越好",
    "",
    "【关键指标：每天至少投 5 家】",
    "投递是概率游戏。按 5% 面试率、20% 面试通过率算，",
    "投 100 家 ≈ 5 个面试 ≈ 1 个 offer。",
    "45 天投不满 60 家，基本等于放弃。",
    "",
    "【面试复盘怎么写才有用】",
    "不要写「问得比较难」这种废话。要写：",
    "  - 具体问了什么（原话）",
    "  - 我当时怎么答的",
    "  - 正确答案应该是什么",
    "第 3-5 次面试时你会发现：问来问去就那些题。",
    "复盘就是把「不会」变成「会」的唯一手段。",
    "",
    "【八股进度怎么用】",
    "按 P0 → P1 → P2 顺序。P0 是面试必问且你截图里被点名的短板。",
    "Go 相关的题不在表里——不会就是不会，硬补性价比太低。",
    "",
    "【一个提醒】",
    "投递优先于学习。学习可以慢慢来，投递窗口不会等你。",
    "不要等「准备好」再投——你永远不会觉得准备好。",
]
for i, line in enumerate(NOTES, start=1):
    c = ws4.cell(row=i, column=1, value=line)
    if i == 1:
        c.font = Font(name="微软雅黑", size=14, bold=True, color="1F3A5F")
    elif line.startswith("【"):
        c.font = Font(name="微软雅黑", size=11, bold=True, color="1F3A5F")
    else:
        c.font = BODY_FONT
    c.alignment = Alignment(wrap_text=False, vertical="center")

wb.save(OUT)
print("saved:", OUT)
print("sheets:", wb.sheetnames)
